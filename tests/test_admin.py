"""Telecaller CRM API — session auth, pipeline updates, stats, DataFrame export."""
import io

import openpyxl


def _seed(client, demo_payload, consult_payload):
    client.post("/api/leads", json=demo_payload)
    client.post("/api/leads", json=consult_payload)


# --- auth gate ------------------------------------------------------------
def test_admin_requires_session(client):
    # No session -> 401 on every admin route.
    assert client.get("/api/admin/leads").status_code == 401
    assert client.get("/api/admin/stats").status_code == 401
    assert client.get("/api/admin/leads/export.csv").status_code == 401
    assert client.get("/api/me").status_code == 401


def test_dashboard_redirects_to_login_when_anonymous(client):
    r = client.get("/dashboard", follow_redirects=False)
    assert r.status_code == 302
    assert r.headers["location"] == "/login"


def test_dashboard_served_when_authenticated(auth_client):
    r = auth_client.get("/dashboard")
    assert r.status_code == 200
    assert b"Telecaller CRM" in r.content


def test_me_returns_user_when_authenticated(auth_client):
    r = auth_client.get("/api/me")
    assert r.status_code == 200
    body = r.json()
    assert body["authenticated"] is True
    assert body["user"]["email"].endswith("@prachalas.com")


# --- pipeline -------------------------------------------------------------
def test_stats_count_by_status(auth_client, demo_payload, consult_payload):
    _seed(auth_client, demo_payload, consult_payload)
    r = auth_client.get("/api/admin/stats")
    assert r.status_code == 200
    stats = r.json()
    assert stats["total"] == 2
    assert stats["demos"] == 1
    assert stats["by_status"]["demo_booked"] == 1
    assert stats["by_status"]["new"] == 1


def test_telecaller_can_work_a_lead(auth_client, consult_payload):
    lead_id = auth_client.post("/api/leads", json=consult_payload).json()["id"]
    r = auth_client.patch(
        f"/api/admin/leads/{lead_id}",
        json={"status": "contacted", "assigned_to": "Anita", "call_notes": "Call back Fri"},
    )
    assert r.status_code == 200
    body = r.json()
    assert body["status"] == "contacted"
    assert body["assigned_to"] == "Anita"
    assert body["call_notes"] == "Call back Fri"


def test_patch_rejects_unknown_status(auth_client, consult_payload):
    lead_id = auth_client.post("/api/leads", json=consult_payload).json()["id"]
    r = auth_client.patch(f"/api/admin/leads/{lead_id}", json={"status": "bogus"})
    assert r.status_code == 422


def test_patch_missing_lead_is_404(auth_client):
    r = auth_client.patch("/api/admin/leads/9999", json={"status": "won"})
    assert r.status_code == 404


def test_status_filter_and_search(auth_client, demo_payload, consult_payload):
    _seed(auth_client, demo_payload, consult_payload)
    only_demo = auth_client.get("/api/admin/leads?status_filter=demo_booked").json()
    assert only_demo["count"] == 1
    assert only_demo["leads"][0]["intent"] == "Schedule a demo"

    found = auth_client.get("/api/admin/leads?q=verma").json()
    assert found["count"] == 1
    assert found["leads"][0]["company"] == "VerMa Realty"


# --- exports --------------------------------------------------------------
def test_csv_export_is_a_dataframe(auth_client, demo_payload, consult_payload):
    _seed(auth_client, demo_payload, consult_payload)
    r = auth_client.get("/api/admin/leads/export.csv")
    assert r.status_code == 200
    assert "text/csv" in r.headers["content-type"]
    text = r.text
    assert "name,company,email" in text.splitlines()[0]
    assert "priya@sunrise.example" in text
    assert "rahul@verma.example" in text


def test_xlsx_export_opens_as_workbook(auth_client, demo_payload):
    auth_client.post("/api/leads", json=demo_payload)
    r = auth_client.get("/api/admin/leads/export.xlsx")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert "email" in header
    assert ws.max_row == 2  # header + 1 data row


# --- logout ---------------------------------------------------------------
def test_logout_clears_session(auth_client):
    assert auth_client.get("/api/me").status_code == 200
    auth_client.get("/auth/logout", follow_redirects=False)
    assert auth_client.get("/api/me").status_code == 401
